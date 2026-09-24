"""
Puts freshly exported Figma frames back into the .pptx.

The deck is authored in Figma and exported one full-bleed PNG per slide, so
"updating the deck" means re-exporting the frames that changed and swapping
those images in — the .pptx is a carrier, not the source. Slides that did not
change keep the bytes they already had.

    python scripts/rebuild-deck.py --slide 2 slides/slide2.png --slide 6 slides/slide6.png

If the .pptx is open in PowerPoint the file is locked; the script then writes
alongside it and says so, rather than failing after doing the work.
"""

import argparse
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "docs" / "NILAMS-SIH2026-Idea-PPT-CupOfTea.pptx"


def replace_artwork(slide, image_path: Path) -> None:
    """Swaps the slide's full-bleed picture for a new export of the frame."""
    picture = next(shape for shape in slide.shapes if shape.shape_type == 13)
    part = picture.part.related_part(picture._element.blip_rId)
    part._blob = image_path.read_bytes()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--slide",
        nargs=2,
        action="append",
        metavar=("NUMBER", "PNG"),
        required=True,
        help="Slide number (1-based) and the PNG exported for it",
    )
    parser.add_argument("--deck", default=str(DECK))
    parser.add_argument("--out", help="Write here instead of over the deck")
    args = parser.parse_args()

    deck = Path(args.deck)
    presentation = Presentation(str(deck))

    for number, png in args.slide:
        index = int(number) - 1
        image = Path(png)
        if not image.exists():
            raise SystemExit(f"missing export: {image}")
        replace_artwork(presentation.slides[index], image)
        print(f"slide {number} <- {image.name} ({image.stat().st_size // 1024} KB)")

    target = Path(args.out) if args.out else deck
    try:
        presentation.save(str(target))
    except PermissionError:
        fallback = deck.with_name(f"{deck.stem}-updated{deck.suffix}")
        presentation.save(str(fallback))
        print(f"\n{deck.name} is open in PowerPoint and locked.")
        print(f"Wrote {fallback.name} instead — close PowerPoint and rename it over the original.")
        return

    print(f"\nwrote {target}")


if __name__ == "__main__":
    main()
